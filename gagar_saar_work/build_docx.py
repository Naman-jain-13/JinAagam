"""Combine parts/partNN.md into one English .docx for Gagar Saar (Gagar Mein Sagar)."""
import glob, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = sys.argv[1] if len(sys.argv) > 1 else "E:/Projects/JinAagam/Gagar_Saar_English.docx"
INPUT = sys.argv[2] if len(sys.argv) > 2 else "E:/Projects/JinAagam/gagar_saar_work/parts/part*.md"
FONT = "Calibri"

META = {
    "pre": "Blessings of",
    "blessing_line": "Param Pujya Vatsalya Varidh Acharya Shiromani 108 Shri Vardhaman Sagar Ji Maharaj",
    "title": "Gagar Mein Sagar",
    "subtitle": "(\u201cAn Ocean in a Pot\u201d)",
    "subtitle2": "A Question-Answer Companion to Jain Agam \u2014 English Edition",
    "compiler": "Compiler: Br. Bhagchand Sethi (presently Muni 108 Shri Prabhavsagar Ji Maharaj)",
    "edition": "Translated into English from the 5th Hindi Edition (2023)",
    "note": "**Translator's note:** This is a paragraph-by-paragraph English translation of the original Hindi question-answer booklet \u201cGagar Mein Sagar,\u201d prepared to make its study material accessible to English readers. The structure \u2014 topics, question numbering, and question-answer format \u2014 is kept identical to the original. Sanskrit/Prakrit technical terms are given in transliteration with an English gloss on first use within each topic. For authoritative citation, always cross-check with the original Hindi text.",
}


def set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(par, text, size=11.5, base_bold=False, color=None):
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            set_font(par.add_run(chunk), size=size, bold=base_bold or i % 2 == 1, color=color)


def add_centered(doc, text, size, color=None, bold=False, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    add_inline(p, text, size=size, base_bold=bold, color=color)
    return p


def add_footer(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    set_font(r, size=10)
    for tag, attr, txt in (("w:fldChar", "begin", None), ("w:instrText", None, " PAGE "), ("w:fldChar", "separate", None), ("w:t", None, "1"), ("w:fldChar", "end", None)):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        if txt:
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        r._r.append(el)


def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = 'TOC \\o "1-1" \\h \\z \\u'
        r._r.append(el)
    t = OxmlElement("w:t")
    t.text = "(Right-click and choose 'Update Field' to build the table of contents)"
    r._r.append(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(end)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    for sname, size, color, space_before, space_after in (
        ("Heading 1", 16, "7A2E00", 14, 8),
        ("Heading 2", 12.5, "1F4E79", 8, 4),
    ):
        st = doc.styles[sname]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)
        st.paragraph_format.keep_with_next = False
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)

    add_footer(doc)

    # --- Cover / title block ---
    for _ in range(3):
        doc.add_paragraph()
    add_centered(doc, META["pre"], 13, color="7A2E00", space_after=4)
    add_centered(doc, META["blessing_line"], 12, color="8B4513", space_after=24)
    add_centered(doc, META["title"], 34, bold=True, color="7A2E00", space_after=6)
    add_centered(doc, META["subtitle"], 15, color="8B4513", space_after=4)
    add_centered(doc, META["subtitle2"], 12, space_after=20)
    add_centered(doc, META["compiler"], 11.5, space_after=4)
    add_centered(doc, META["edition"], 11, color="444444", space_after=30)
    p = doc.add_paragraph()
    add_inline(p, META["note"], size=10.5, color="444444")
    doc.add_page_break()

    doc.add_heading("Contents", level=1)
    add_toc(doc)
    doc.add_page_break()

    files = sorted(glob.glob(INPUT))
    for f in files:
        lines = open(f, encoding="utf-8").read().splitlines()
        i = 0
        while i < len(lines):
            ln = lines[i]
            s = ln.strip()
            if s.startswith("## "):
                h = doc.add_heading(level=2)
                add_inline(h, s[3:], size=12.5, color="1F4E79")
            elif s.startswith("# "):
                h = doc.add_heading(level=1)
                add_inline(h, s[2:], size=16, color="7A2E00")
            elif s.startswith("**Question"):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, s, size=11.5)
            elif s.startswith("Answer"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(4)
                add_inline(p, s, size=11.5, color="2B2B2B")
            elif s == "---" or not s:
                pass
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                add_inline(p, s)
            i += 1
    doc.save(OUT)
    print("saved", OUT, "from", len(files), "parts")


main()
