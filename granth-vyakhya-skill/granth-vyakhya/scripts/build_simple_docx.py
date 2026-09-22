"""Simple builder for translation-only work (question-answer booklets, plain prose translations) — no 6-part § structure.

Usage:  python build_simple_docx.py <work_dir>
  <work_dir>/meta.json  : {"title", "subtitle", "lines": [...], "note", "out", "font", "toc_levels": "1-1"}
  <work_dir>/parts/*.md : "# " Heading 1, "## " Heading 2, "**Question N.** ..." / "Answer: ..." pairs,
                          "- " bullets, "| a | b |" tables, plain paragraphs. Files are concatenated in name order.
Then run finalize_word.ps1 on the .docx to fill the TOC page numbers and export the PDF.
"""
import glob, json, os, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

WORK = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else ".")
META = json.load(open(os.path.join(WORK, "meta.json"), encoding="utf-8"))
FONT = META.get("font", "Nirmala UI")
OUT = os.environ.get("OUT") or os.path.abspath(os.path.join(WORK, META.get("out", "../Translation.docx")))


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)


def add_inline(par, text, size=11.5, base_bold=False, color=None):
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            set_font(par.add_run(chunk), size=size, bold=base_bold or i % 2 == 1, color=color)


def centered(doc, text, size, color=None, bold=False, after=6):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after); add_inline(p, text, size, bold, color); return p


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.fullmatch(r"\|?[\s:|-]+\|?", r.strip())]
    n = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=n); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(cells):
        for j in range(n):
            c = t.cell(i, j); add_inline(c.paragraphs[0], r[j] if j < len(r) else "", 10.5, i == 0)
            if i == 0:
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F3E3C3"); c._element.get_or_add_tcPr().append(shd)
    doc.add_paragraph()


def add_field(run, instr):
    for tag, attr, txt in (("w:fldChar", "begin", None), ("w:instrText", None, instr), ("w:fldChar", "separate", None), ("w:t", None, "…"), ("w:fldChar", "end", None)):
        el = OxmlElement(tag)
        if attr: el.set(qn("w:fldCharType"), attr)
        if txt: el.set(qn("xml:space"), "preserve"); el.text = txt
        run._r.append(el)


doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.5)
for name, size, color in (("Heading 1", 16, "7A2E00"), ("Heading 2", 12.5, "1F4E79")):
    st = doc.styles[name]; st.font.name = FONT; st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.keep_with_next = True  # never leave a heading alone at a page bottom
    rpr = st.element.get_or_add_rPr(); rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    rpr.append(rf)

# footer page number (set once on section 0 only — later linked sections share it)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run(); set_font(fr, 10); add_field(fr, " PAGE ")

for _ in range(3): doc.add_paragraph()
for line in META.get("pre_lines", []): centered(doc, line, 12, "8B4513")
centered(doc, META["title"], 32, "7A2E00", True, 8)
if META.get("subtitle"): centered(doc, META["subtitle"], 14, "8B4513", after=16)
for line in META.get("lines", []): centered(doc, line, 11.5)
if META.get("note"):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(24); add_inline(p, META["note"], 10.5, color="444444")
doc.add_page_break()
doc.add_heading(META.get("toc_title", "Contents"), level=1)
p = doc.add_paragraph(); r = p.add_run(); add_field(r, f'TOC \\o "{META.get("toc_levels", "1-1")}" \\h \\z \\u')
doc.add_page_break()

for f in sorted(glob.glob(os.path.join(WORK, "parts", "*.md"))):
    lines = open(f, encoding="utf-8").read().splitlines(); i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"): rows.append(lines[i]); i += 1
            add_table(doc, rows); continue
        if s.startswith("## "): add_inline(doc.add_heading(level=2), s[3:], 12.5, color="1F4E79")
        elif s.startswith("# "): add_inline(doc.add_heading(level=1), s[2:], 16, color="7A2E00")
        elif s.startswith("**Question") or s.startswith("**प्रश्न"):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2); add_inline(p, s)
        elif s.startswith("Answer") or s.startswith("उत्तर"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(4); add_inline(p, s, color="2B2B2B")
        elif re.match(r"^[-*•] ", s): add_inline(doc.add_paragraph(style="List Bullet"), s[2:])
        elif s and s != "---":
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); add_inline(p, s)
        i += 1

doc.save(OUT)
print("saved", OUT)
