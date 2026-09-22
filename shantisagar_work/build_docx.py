"""Combine parts/partNN.md into one English .docx — Shantisagar Ji (English Edition).
Markdown subset: # H1, ## H2, **bold**, *italic*, | tables |, **Question N —** / Answer — lines, > quote, - bullets.
Then run finalize_word.ps1 (OUT env) to update TOC + export PDF."""
import glob, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = sys.argv[1] if len(sys.argv) > 1 else "E:/Projects/JinAagam/Shantisagar_Ji_English.docx"
INPUT = sys.argv[2] if len(sys.argv) > 2 else "E:/Projects/JinAagam/shantisagar_work/parts/part*.md"
FONT = "Calibri"

META = {
    "pre": "Ever to Be Remembered at Dawn — the First Acharya of the 20th Century",
    "title": "Charitra Chakravarti Acharya 108 Shri Shantisagar Ji Maharaj (Dakshin)",
    "subtitle": "Life-Sketch \u00b7 Chaturmas and Disciple Records \u00b7 Question-Answer Quiz \u00b7 Pooja \u00b7 Aarti \u00b7 Chalisa \u00b7 The Twelve Bhavanas",
    "subtitle2": "English Edition",
    "edition": "Translated into English from the 40-page Hindi booklet \u201cShantisagar Ji\u201d",
    "note": "**Translator\u2019s note:** This is a page-by-page English translation of the Hindi booklet on Charitra Chakravarti Acharya Shri Shantisagar Ji Maharaj (Dakshin), prepared to make it accessible to English readers. The structure of the original \u2014 the life-sketch, the chaturmas and disciple tables with their sources, the 192-question quiz, the Pooja, Aarti, Chalisa and the Twelve Bhavanas \u2014 is kept identical. Devotional verses are given first in Roman transliteration (so they can be recited as in the original) and then in English meaning; the Sanskrit offering-mantras are transliterated. Sanskrit/Prakrit technical terms appear in transliteration with an English gloss on first use. Names of places and persons follow the spellings in the original. For authoritative citation, always cross-check with the original Hindi text.",
}

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(par, text, size=11.5, base_bold=False, base_italic=False, color=None):
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            set_font(par.add_run(chunk[2:-2]), size=size, bold=True, italic=base_italic, color=color)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            set_font(par.add_run(chunk[1:-1]), size=size, bold=base_bold, italic=True, color=color)
        else:
            set_font(par.add_run(chunk), size=size, bold=base_bold, italic=base_italic, color=color)


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_table(doc, rows):
    rows = [r for r in rows if not re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", r.strip())]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    ncol = max(len(r) for r in cells)
    header = any(cells[0]) if cells else False
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    total = 17.0  # cm usable width
    first = cells[0][0] if cells and cells[0] else ""
    if first in ("No.", "Nos."):
        w0 = 1.2 if ncol > 2 else 2.2
    elif ncol == 2:
        w0 = 4.6
    else:
        w0 = total / ncol
    rest = (total - w0) / (ncol - 1) if ncol > 1 else 0
    widths = [w0] + [rest] * (ncol - 1)
    for i, r in enumerate(cells):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            add_inline(p, r[j] if j < len(r) else "", size=10, base_bold=(i == 0 and header))
            if i == 0 and header:
                shade(cell, "F3E3C3")
    doc.add_paragraph()


def add_centered(doc, text, size, color=None, bold=False, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    add_inline(p, text, size=size, base_bold=bold, color=color)
    return p


def add_footer(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    set_font(r, size=10)
    for tag, attr, txt in (("w:fldChar", "begin", None), ("w:instrText", None, " PAGE "), ("w:fldChar", "separate", None), ("w:t", None, "1"), ("w:fldChar", "end", None)):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        if txt:
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        r._r.append(el)


def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = 'TOC \\o "1-2" \\h \\z \\u'
        r._r.append(el)
    t = OxmlElement("w:t")
    t.text = "(Right-click and choose 'Update Field' to build the table of contents)"
    r._r.append(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(end)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    for sname, size, color, space_before, space_after in (
        ("Heading 1", 16, "7A2E00", 14, 8),
        ("Heading 2", 12.5, "1F4E79", 10, 4),
    ):
        st = doc.styles[sname]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)
        st.paragraph_format.keep_with_next = True
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)

    add_footer(doc)

    # --- Cover / title block ---
    for _ in range(3):
        doc.add_paragraph()
    add_centered(doc, META["pre"], 12, color="8B4513", space_after=18)
    add_centered(doc, META["title"], 24, bold=True, color="7A2E00", space_after=8)
    add_centered(doc, META["subtitle"], 12.5, color="8B4513", space_after=4)
    add_centered(doc, META["subtitle2"], 13, bold=True, space_after=20)
    add_centered(doc, META["edition"], 11, color="444444", space_after=30)
    p = doc.add_paragraph()
    add_inline(p, META["note"], size=10.5, color="444444")
    doc.add_page_break()

    doc.add_heading("Contents", level=1)
    add_toc(doc)
    doc.add_page_break()

    files = sorted(glob.glob(INPUT))
    first_h1 = True
    for f in files:
        lines = open(f, encoding="utf-8").read().splitlines()
        i = 0
        while i < len(lines):
            ln = lines[i]
            s = ln.strip()
            if s.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append(lines[i])
                    i += 1
                add_table(doc, rows)
                continue
            if s.startswith("## "):
                h = doc.add_heading(level=2)
                add_inline(h, s[3:], size=12.5, color="1F4E79")
            elif s.startswith("# "):
                if not first_h1:
                    doc.add_page_break()
                first_h1 = False
                h = doc.add_heading(level=1)
                add_inline(h, s[2:], size=16, color="7A2E00")
            elif s.startswith("**Question"):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, s, size=11.5)
            elif s.startswith("Answer"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(4)
                add_inline(p, s, size=11.5, color="2B2B2B")
            elif s.startswith("> "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                add_inline(p, s[2:], color="444444")
            elif re.match(r"^\s*[-*] ", ln) and not s.startswith("* * *"):
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, re.sub(r"^\s*[-*] ", "", ln))
            elif s == "* * *":
                add_centered(doc, "\u2767 \u2766 \u2767", 12, color="D9C9A8", space_before=6, space_after=6)
            elif s == "---" or not s:
                pass
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                add_inline(p, s)
            i += 1
    doc.save(OUT)
    print("saved", OUT, "from", len(files), "parts")


main()
