# -*- coding: utf-8 -*-
"""निष्कर्षित मूल संस्कृत शास्त्र -> .docx
   प्रयोग: python build_mool_docx.py
   इसके बाद finalize_word.ps1 से TOC के पृष्ठांक भरें और PDF बनाएँ।
"""
import glob, io, json, os, re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

WORK = os.path.abspath(".")
META = json.load(io.open(os.path.join(WORK, "mool_book", "meta.json"), encoding="utf-8"))
FONT = META.get("font", "Nirmala UI")
OUT  = os.path.abspath(os.path.join(WORK, "mool_book", META["out"]))

def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)

def para(text, size=11.5, align=None, color=None, bold=False, before=0, after=5,
         indent=None, line=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent is not None: pf.left_indent = Cm(indent)
    if line is not None: pf.line_spacing = line
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p

def add_field(run, instr):
    for tag, attr, txt in (("w:fldChar","begin",None),("w:instrText",None,instr),
                           ("w:fldChar","separate",None),("w:t",None,"…"),("w:fldChar","end",None)):
        el = OxmlElement(tag)
        if attr: el.set(qn("w:fldCharType"), attr)
        if txt: el.set(qn("xml:space"), "preserve"); el.text = txt
        run._r.append(el)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(1.8)

for name, size, color in (("Heading 1", 15, "7A2E00"), ("Heading 2", 12, "1F4E79")):
    st = doc.styles[name]; st.font.name = FONT; st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.keep_with_next = True
    rpr = st.element.get_or_add_rPr(); rf = OxmlElement("w:rFonts")
    for a in ("w:ascii","w:hAnsi","w:cs","w:eastAsia"): rf.set(qn(a), FONT)
    rpr.append(rf)

fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run(); set_font(fr, 10); add_field(fr, " PAGE ")

C = WD_ALIGN_PARAGRAPH.CENTER; J = WD_ALIGN_PARAGRAPH.JUSTIFY

# ---------- आवरण ----------
for _ in range(3): doc.add_paragraph()
para(META["author_line"], 13, C, "8B4513", after=10)
para(META["title"], 34, C, "7A2E00", True, after=8)
para(META["subtitle"], 13.5, C, "8B4513", after=20)
for l in META.get("lines", []): para(l, 11.5, C, after=3)
para(META["basis"], 10, J, "444444", before=20, after=8)
para(META["note"], 10, J, "8B0000", after=8)
para(META["disclaimer"], 9.5, J, "444444")
doc.add_page_break()

# ---------- विषय-सूची ----------
para("विषय-सूची", 15, color="7A2E00", bold=True, after=8)   # Heading 1 नहीं — वरना TOC में स्वयं आ जाता है
p = doc.add_paragraph(); r = p.add_run(); add_field(r, "TOC " + chr(92) + chr(111) + " " + chr(34) + "1-1" + chr(34) + " " + chr(92) + "h " + chr(92) + "z " + chr(92) + "u")
doc.add_page_break()

# ---------- पाठ ----------
VERSE = re.compile(r'॥\s*[०-९0-9]+\s*॥\s*$')

def mark_verses(lines):
    """कारिका के सभी चरण चिह्नित करो — केवल ॥N॥ वाली अन्तिम पंक्ति नहीं।
       पद्य 2 या 4 चरणों का होता है; पीछे की ओर चलकर सब समेटो।"""
    v = [False] * len(lines)
    for i, l in enumerate(lines):
        if VERSE.search(l.strip()):
            v[i] = True
            j = i - 1
            steps = 0
            while j >= 0 and steps < 3:
                t = lines[j].strip()
                if not t: break
                if t.startswith(('#', '[', '{', '§', '<!--')) or t.startswith('###'): break
                if re.match(r'^\[[०-९\d]+\]', t): break
                if not t.endswith('।'): break
                v[j] = True; j -= 1; steps += 1
    return v

mode = None
for f in sorted(glob.glob(os.path.join(WORK, "mool_book", "parts", "*.md"))):
    lines = io.open(f, encoding="utf-8").read().splitlines()
    isverse = mark_verses(lines)
    titleblk = True
    i = 0
    while i < len(lines):
        s = lines[i].strip(); was_verse = isverse[i]; i += 1
        if not s: continue
        s = re.sub(r'^\(\s*पूर्व\s*§[^)]*\)\s*', '', s).strip()
        if not s: continue
        if titleblk and not s.startswith(("[", "#", "§", "{", "<!--")):
            para(s, 16 if "परीक्षा" in s and len(s) < 22 else 12,
                 C, "7A2E00", True, before=3, after=3); continue
        m = re.match(r'<!--PG (\d+)-->', s)
        if m:
            para("[ मुद्रित पृष्ठ %s ]" % m.group(1), 7.5, WD_ALIGN_PARAGRAPH.RIGHT,
                 "999999", before=2, after=0)
            continue
        if s.startswith("# "):
            doc.add_page_break()
            hh = doc.add_heading(level=1); hh.alignment = C
            set_font(hh.add_run(s[2:]), 15, bold=True, color="7A2E00")
            mode = None; continue
        if s.startswith("### "):
            mode = s[4:].strip()
            para(mode, 9, color="7A2E00", bold=True, before=6, after=1); continue
        if s.startswith("[") and s.endswith("]"):
            titleblk = False
            para(s, 11, C, "1F4E79", True, before=9, after=4); mode = None; continue
        if re.match(r'^\{\d+\}', s) or re.match(r'^\[[०-९\d]+\]', s):
            para(s, 7.5, indent=0.8, after=0, line=0.95, color="333333"); continue
        if s.startswith("§"):
            mode = None
            hd = doc.add_heading(level=2)
            mm = re.match(r'^(§\s*\d+\.?)\s*(.*)$', s)
            lbl, rest = (mm.group(1), mm.group(2)) if mm else (s, "")
            set_font(hd.add_run(lbl), 11, bold=True, color="1F4E79")
            hd.paragraph_format.space_before = Pt(8); hd.paragraph_format.space_after = Pt(2)
            if rest: para(rest, 11.5, J, after=5, line=1.15)
            continue
        if was_verse:
            para(s, 12.5, C, "5B2C00", True, before=2, after=2); continue
        para(s, 11.5, J, after=5, line=1.15)

doc.save(OUT)
print("saved", OUT)
