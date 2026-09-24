"""Book-format build (5.5in x 10.5in, like 'Gagar Me Sagar 23-9-2025.pdf'):
prose/tables single column; the Question-Answer quiz in two columns with a rule; footer '( N )'.
Then: finalize_word.ps1 with $env:OUT to update TOC + export PDF."""
import glob, re, sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Inches

HERE = __import__("pathlib").Path(__file__).resolve().parent
OUT = sys.argv[1] if len(sys.argv) > 1 else str(HERE / "vyakhya" / "Shantisagar_Ji_English_Print.docx")
INPUT = sys.argv[2] if len(sys.argv) > 2 else str(HERE / "parts/part*.md")
FONT = "Cambria"
BODY = 10
PAGE_W, PAGE_H = Inches(5.5), Inches(10.5)
M_LR, M_TB = Inches(0.6), Inches(0.85)
USABLE_CM = (5.5 - 1.2) * 2.54  # 10.9 cm

META = {
    "pre": "Ever to Be Remembered at Dawn \u2014 the First Acharya of the 20th Century",
    "title": "Charitra Chakravarti Acharya 108 Shri Shantisagar Ji Maharaj (Dakshin)",
    "subtitle": "Life-Sketch \u00b7 Chaturmas and Disciple Records \u00b7 Question-Answer Quiz \u00b7 Pooja \u00b7 Aarti \u00b7 Chalisa \u00b7 The Twelve Bhavanas",
    "subtitle2": "English Edition",
    "edition": "Translated into English from the 40-page Hindi booklet \u201cShantisagar Ji\u201d",
    "note": "**Translator\u2019s note:** This is a page-by-page English translation of the Hindi booklet on Charitra Chakravarti Acharya Shri Shantisagar Ji Maharaj (Dakshin). The structure of the original \u2014 the life-sketch, the chaturmas and disciple tables with their sources, the 192-question quiz, the Pooja, Aarti, Chalisa and the Twelve Bhavanas \u2014 is kept identical. Devotional verses are given first in Roman transliteration (so they can be recited as in the original) and then in English meaning; the Sanskrit offering-mantras are transliterated. Names of places and persons follow the spellings in the original. For authoritative citation, always cross-check with the original Hindi text.",
}
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)


def add_inline(par, text, size=BODY, base_bold=False, base_italic=False, color=None):
    for chunk in INLINE_RE.split(text):
        if not chunk: continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            set_font(par.add_run(chunk[2:-2]), size=size, bold=True, italic=base_italic, color=color)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            set_font(par.add_run(chunk[1:-1]), size=size, bold=base_bold, italic=True, color=color)
        else:
            set_font(par.add_run(chunk), size=size, bold=base_bold, italic=base_italic, color=color)


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_table(doc, rows):
    rows = [r for r in rows if not re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", r.strip())]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    ncol = max(len(r) for r in cells)
    header = any(cells[0]) if cells else False
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    first = cells[0][0] if cells and cells[0] else ""
    if first in ("No.", "Nos."):
        w0 = 0.9 if ncol > 2 else 1.6
    elif ncol == 2:
        w0 = 3.4
    else:
        w0 = USABLE_CM / ncol
    rest = (USABLE_CM - w0) / (ncol - 1) if ncol > 1 else 0
    widths = [w0] + [rest] * (ncol - 1)
    fs = 7.5 if ncol >= 5 else 8.5
    for i, r in enumerate(cells):
        for j in range(ncol):
            cell = t.cell(i, j); cell.width = Cm(widths[j])
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_inline(p, r[j] if j < len(r) else "", size=fs, base_bold=(i == 0 and header))
            if i == 0 and header: shade(cell, "F3E3C3")
    doc.add_paragraph()


def add_centered(doc, text, size, color=None, bold=False, space_before=0, space_after=0):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(space_after)
    add_inline(p, text, size=size, base_bold=bold, color=color)
    return p


def add_footer(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("( "); set_font(r, size=9)
    r = p.add_run(); set_font(r, size=9)
    for tag, attr, txt in (("w:fldChar", "begin", None), ("w:instrText", None, " PAGE "), ("w:fldChar", "separate", None), ("w:t", None, "1"), ("w:fldChar", "end", None)):
        el = OxmlElement(tag)
        if attr: el.set(qn("w:fldCharType"), attr)
        if txt: el.set(qn("xml:space"), "preserve"); el.text = txt
        r._r.append(el)
    r = p.add_run(" )"); set_font(r, size=9)


def add_toc(doc):
    p = doc.add_paragraph(); r = p.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate")):
        el = OxmlElement(tag)
        if attr: el.set(qn("w:fldCharType"), attr)
        else: el.set(qn("xml:space"), "preserve"); el.text = 'TOC \\o "1-1" \\h \\z \\u'
        r._r.append(el)
    t = OxmlElement("w:t"); t.text = "(Update Field to build the contents)"; r._r.append(t)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); r._r.append(end)


def setup_section(sec, cols=1):
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.left_margin = sec.right_margin = M_LR
    sec.top_margin = sec.bottom_margin = M_TB
    sec.footer_distance = Inches(0.45)
    sectPr = sec._sectPr
    for old in sectPr.findall(qn("w:cols")): sectPr.remove(old)
    c = OxmlElement("w:cols"); c.set(qn("w:num"), str(cols))
    if cols > 1:
        c.set(qn("w:sep"), "1"); c.set(qn("w:space"), "280")
    # w:cols must precede w:docGrid / after pgMar; append before docGrid if present
    dg = sectPr.find(qn("w:docGrid"))
    if dg is not None: dg.addprevious(c)
    else: sectPr.append(c)


def new_section(doc, cols):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(sec, cols)
    return sec


def main():
    doc = Document()
    setup_section(doc.sections[0], 1)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(BODY)
    for sname, size, color, sb, sa in (("Heading 1", 13, "7A2E00", 10, 6), ("Heading 2", 11, "1F4E79", 8, 3)):
        s = doc.styles[sname]; s.font.name = FONT; s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(sb); s.paragraph_format.space_after = Pt(sa); s.paragraph_format.keep_with_next = True
        rpr = s.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    add_footer(doc)

    # cover
    for _ in range(4): doc.add_paragraph()
    add_centered(doc, "\u0965 Shri Shanti-Vira-Shiva-Dharmajita-Vardhamana Suribhyo Namah \u0965", 9, color="8B4513", space_after=14)
    add_centered(doc, META["pre"], 10, color="8B4513", space_after=14)
    add_centered(doc, META["title"], 19, bold=True, color="7A2E00", space_after=8)
    add_centered(doc, META["subtitle"], 10, color="8B4513", space_after=4)
    add_centered(doc, META["subtitle2"], 11.5, bold=True, space_after=16)
    add_centered(doc, META["edition"], 9.5, color="444444", space_after=26)
    p = doc.add_paragraph(); add_inline(p, META["note"], size=8.5, color="444444")
    doc.add_page_break()
    doc.add_heading("Contents", level=1); add_toc(doc)

    files = sorted(glob.glob(INPUT))
    in_two_col = False
    first_h1 = True
    for f in files:
        lines = open(f, encoding="utf-8").read().splitlines()
        i = 0
        while i < len(lines):
            ln = lines[i]; s = ln.strip()
            if s.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append(lines[i]); i += 1
                add_table(doc, rows); continue
            if s.startswith("# "):
                title = s[2:]
                want_two = title.startswith("Charitra Chakravarti Shri 108 Shantisagar Question-Answer Quiz")
                if want_two != in_two_col:
                    new_section(doc, 2 if want_two else 1); in_two_col = want_two
                else:
                    doc.add_page_break()
                first_h1 = False
                h = doc.add_heading(level=1); add_inline(h, title, size=13, color="7A2E00")
            elif s.startswith("## "):
                h = doc.add_heading(level=2); add_inline(h, s[3:], size=11, color="1F4E79")
            elif s.startswith("**Question"):
                p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.keep_with_next = True
                add_inline(p, s.replace("**Question ", "**Q. ", 1), size=BODY)
            elif s.startswith("Answer"):
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline(p, "Ans." + s[len("Answer"):], size=BODY, color="222222")
            elif s.startswith("> "):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); add_inline(p, s[2:], color="444444")
            elif re.match(r"^\s*[-*] ", ln) and s != "* * *":
                p = doc.add_paragraph(style="List Bullet"); add_inline(p, re.sub(r"^\s*[-*] ", "", ln))
            elif s == "* * *":
                add_centered(doc, "\u2767 \u2766 \u2767", 11, color="D9C9A8", space_before=4, space_after=4)
            elif s == "---" or not s:
                pass
            else:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
                if s.startswith("*") and s.endswith("*") and not s.startswith("**"):
                    p.paragraph_format.space_after = Pt(1)  # verse line
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline(p, s)
            i += 1
    doc.save(OUT); print("saved", OUT, "from", len(files), "parts")


main()
